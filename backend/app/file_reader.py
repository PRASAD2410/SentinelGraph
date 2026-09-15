from io import BytesIO
from pathlib import Path
import csv
from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook

ALLOWED={'.txt','.csv','.pdf','.docx','.xlsx'}
def read_upload(filename: str, content: bytes) -> str:
    suffix=Path(filename).suffix.lower()
    if suffix not in ALLOWED: raise ValueError('Supported formats: PDF, DOCX, TXT, CSV, XLSX.')
    if suffix=='.txt': return content.decode('utf-8',errors='replace')
    if suffix=='.pdf': return '\n'.join(page.extract_text() or '' for page in PdfReader(BytesIO(content)).pages)
    if suffix=='.docx': return '\n'.join(p.text for p in Document(BytesIO(content)).paragraphs)
    if suffix=='.csv': return '\n'.join(' | '.join(row) for row in csv.reader(content.decode('utf-8',errors='replace').splitlines()))
    workbook=load_workbook(BytesIO(content),read_only=True,data_only=True)
    return '\n'.join(' | '.join(str(v or '') for v in row) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True))
